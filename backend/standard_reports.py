"""Offline, admin-only Excel/PDF exports and print-friendly results."""
from io import BytesIO
from pathlib import Path
from html import escape
import os
from fastapi import Depends, Query
from fastapi.responses import StreamingResponse, HTMLResponse
from backend.app import app, admin_only
from backend.standard_admin_results import admin_standard_results

FIELDS = ("full_name","username","grade","exam_title","exam_type","score","maximum","submitted_at")
HEADERS = ("Student / الطالب","Username","Grade / الصف","Exam / الامتحان","Type","Score","Max","Date / التاريخ")

def data_for(grade, user):
    return admin_standard_results(grade=grade,limit=50000,user=user)

def excel_safe(value):
    value=str(value if value is not None else "")
    return "'"+value if value.lstrip().startswith(("=","+","-","@")) else value

@app.get("/api/admin/standard/export.xlsx")
def export_excel(grade: str | None = Query(default=None,max_length=80), user=Depends(admin_only)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins
    data=data_for(grade,user)
    wb=Workbook();ws=wb.active;ws.title="Student Results"
    ws.merge_cells("A1:H2");ws["A1"]="THE SPECIAL ONE | STUDENT RESULTS"
    ws["A1"].font=Font(name="Calibri",size=20,bold=True,color="FFFFFF")
    ws["A1"].fill=PatternFill("solid",fgColor="142B49")
    ws["A1"].alignment=Alignment(vertical="center",horizontal="center")
    ws.merge_cells("A3:H3");ws["A3"]=f"Grade: {grade or 'All'}    |    Students: {data['students']}    |    Teacher: {data['teacher_attempts']}    |    Practice: {data['practice_attempts']}    |    Total: {data['total']}"
    ws["A3"].font=Font(name="Calibri",bold=True,color="24476C")
    for i,title in enumerate(HEADERS,1):
        cell=ws.cell(5,i,title);cell.font=Font(name="Calibri",bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor="2563A0");cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    for n,item in enumerate(data["results"],6):
        for i,key in enumerate(FIELDS,1):
            value=item.get(key,"")
            if key in ("score","maximum"): value=int(value)
            else: value=excel_safe(value)
            cell=ws.cell(n,i,value)
            cell.font=Font(name="Calibri",size=11,color="182B45")
            cell.fill=PatternFill("solid",fgColor="F0F5FA" if n%2==0 else "FFFFFF")
            cell.alignment=Alignment(vertical="center",wrap_text=True)
        ws.row_dimensions[n].height=24
    for i,w in enumerate((30,22,18,35,18,12,12,25),1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[5].height=30;ws.freeze_panes="D6";ws.auto_filter.ref=f"A5:H{max(5,5+len(data['results']))}"
    ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation="landscape";ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows="1:5"
    ws.page_margins=PageMargins(left=.25,right=.25,top=.4,bottom=.4,header=.2,footer=.2)
    ws.oddFooter.center.text="THE SPECIAL ONE • Page &P / &N"
    buf=BytesIO();wb.save(buf);buf.seek(0)
    return StreamingResponse(buf,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":'attachment; filename="THE_SPECIAL_ONE_Results.xlsx"',"Cache-Control":"no-store"})

def pdf_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    candidates=[Path(os.environ.get("WINDIR","C:/Windows"))/"Fonts"/"arial.ttf",Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")]
    for path in candidates:
        if path.is_file():
            pdfmetrics.registerFont(TTFont("TSOUnicode",str(path)))
            return "TSOUnicode"
    return "Helvetica"

def bidi_text(value):
    import arabic_reshaper
    from bidi.algorithm import get_display
    return get_display(arabic_reshaper.reshape(str(value)))

@app.get("/api/admin/standard/export.pdf")
def export_pdf(grade: str | None = Query(default=None,max_length=80), user=Depends(admin_only)):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4,landscape
    data=data_for(grade,user);buf=BytesIO();font=pdf_font()
    doc=SimpleDocTemplate(buf,pagesize=landscape(A4),leftMargin=25,rightMargin=25,topMargin=28,bottomMargin=28)
    title=ParagraphStyle("title",fontName=font,fontSize=17,leading=24,textColor=colors.HexColor("#142B49"),alignment=TA_CENTER)
    normal=ParagraphStyle("cell",fontName=font,fontSize=8,leading=11)
    story=[Paragraph("THE SPECIAL ONE | STUDENT RESULTS",title),Spacer(1,10)]
    story.append(Paragraph(escape(f"Grade: {grade or 'All'}  |  Students: {data['students']}  |  Teacher: {data['teacher_attempts']}  |  Practice: {data['practice_attempts']}  |  Total: {data['total']}"),normal))
    story.append(Spacer(1,12))
    columns=("Student","Grade","Exam","Type","Score","Max","Date")
    widths=(150,85,185,65,42,42,160)
    cells=[[Paragraph(escape(x),normal) for x in columns]]
    for item in data["results"]:
        values=(item["full_name"],item["grade"],item["exam_title"],item["exam_type"],item["score"],item["maximum"],item["submitted_at"][:16].replace("T"," "))
        cells.append([Paragraph(escape(bidi_text(x) if isinstance(x,str) else str(x)),normal) for x in values])
    table=Table(cells,colWidths=widths,repeatRows=1,hAlign="CENTER")
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#DCEAF5")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F3F7FB")]),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("BOTTOMPADDING",(0,0),(-1,-1),7),("TOPPADDING",(0,0),(-1,-1),7),("LINEBELOW",(0,0),(-1,0),.8,colors.HexColor("#24476C"))]))
    story.append(table)
    def footer(canvas,document):
        canvas.saveState();canvas.setFont(font,8);canvas.drawString(26,15,"THE SPECIAL ONE • OFFLINE STUDENT MANAGEMENT");canvas.drawRightString(landscape(A4)[0]-26,15,f"Page {document.page}");canvas.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer);buf.seek(0)
    return StreamingResponse(buf,media_type="application/pdf",headers={"Content-Disposition":'attachment; filename="THE_SPECIAL_ONE_Results.pdf"',"Cache-Control":"no-store"})

@app.get("/api/admin/standard/print",response_class=HTMLResponse)
def print_results(grade: str | None = Query(default=None,max_length=80), user=Depends(admin_only)):
    data=data_for(grade,user)
    rows="".join("<tr>"+"".join(f"<td>{escape(str(item.get(key,'')))}</td>" for key in FIELDS)+"</tr>" for item in data["results"])
    head="".join(f"<th>{escape(label)}</th>" for label in HEADERS)
    return HTMLResponse("""<!doctype html><html lang="en"><head><meta charset="utf-8"><title>THE SPECIAL ONE — Student Results</title><style>
    @page{size:A4 landscape;margin:12mm}body{font-family:Arial,sans-serif;color:#18324b;margin:22px}header{border-bottom:3px solid #235b92;padding-bottom:12px;margin-bottom:16px}h1{margin:0;font-size:23px}p{font-size:12px}table{width:100%;border-collapse:collapse;font-size:10px}th{background:#214d77;color:white;text-align:left}td,th{padding:8px;border-bottom:1px solid #d8e2ec;overflow-wrap:anywhere}tbody tr:nth-child(even){background:#f2f6fa}thead{display:table-header-group}tr{break-inside:avoid}button{padding:12px 20px;background:#235b92;color:white;border:0;border-radius:7px;margin-bottom:16px;cursor:pointer}@media print{button{display:none}body{margin:0}th{-webkit-print-color-adjust:exact;print-color-adjust:exact}tr{-webkit-print-color-adjust:exact;print-color-adjust:exact}}</style></head><body><button onclick="window.print()">Print / طباعة</button><header><h1>THE SPECIAL ONE</h1><p>STUDENT RESULTS • Grade: """+escape(grade or "All")+f""" • Students: {data['students']} • Teacher exams: {data['teacher_attempts']} • Practice: {data['practice_attempts']} • Total: {data['total']}</p></header><table><thead><tr>"""+head+"</tr></thead><tbody>"+rows+"</tbody></table></body></html>",headers={"Cache-Control":"no-store"})
